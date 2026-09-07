from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

d = Document()
st = d.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)
d.add_heading('Freight Rate Prediction - Solution Report', 0)
d.add_paragraph('Spotter.AI | Machine Learning Engineer Assessment').italic = True
d.add_paragraph('This report describes what I did, why I did it, and what came out. It covers the data split and validation approach, and includes the December prediction chart from score.py as required.')

def h(x):
    d.add_heading(x, 1)

def h2(x):
    d.add_heading(x, 2)

def p(x):
    d.add_paragraph(x)

def bullets(items):
    for i in items:
        d.add_paragraph(i, style='List Bullet')

h('1. Objective in plain words')
p('Each row is one truck load someone wants to move. The job is to predict the price (posted_rate, in dollars) before the truck is booked. I trained on 48,000 old loads from Jan to Oct 2025, then predicted 12,000 new loads from Nov to Dec 2025, plus a fixed 31-day December lane (Lexington to Fort Wayne) that tests whether the model understands time when nothing else changes.')

h('2. Dataset')
bullets([
    'Train (Data/train-test.csv): 48,000 rows, 14 columns, dates 2025-01-01 to 2025-10-31. Has the target posted_rate.',
    'Validation (Data/validation.csv): 12,000 rows, 13 columns (no target), dates 2025-11-01 to 2025-12-31. IDs TE-000001 to TE-012000.',
    'Template (Data/validation-predictions-template.csv): 12,000 IDs to fill.',
    'December probe (Data/december-chart-inputs.csv): 31 rows, Dec 1-31, fixed lane Lexington to Fort Wayne, 360 miles, Dry Van, 32,000 lb. Only date changes.',
    'Target stats: mean $2,373.98, median $2,030.76, std $1,486.49, min $57.22, max $25,533, skew 1.90. Strongly right-skewed.',
    'Columns: load_id (ID only, never a feature), pickup/delivery cities (64 each, 4,014 routes), pickup/delivery lat-lon, distance, equipment (Dry Van / Reefer / Flatbed), weight, date, market_index, quote_signal.',
])

h('3. Data quality issues and how I fixed them')
p('I checked missing values, duplicates, negatives, and impossible values before modelling anything. No duplicate rows or IDs in either file. Coordinates, distance, dates were complete.')
bullets([
    'Negative weights: 292 in train, 145 in validation. These looked like sign-entry mistakes (magnitude normal, e.g. -30,000 lb), so I took absolute values instead of deleting them.',
    'Missing weight: 300 train, 165 validation. Fixed with equipment-specific median (Dry Van / Reefer / Flatbed separately), because truck types carry different typical weights. Fallback to global median if a group were empty.',
    'Missing market_index: 374 train, 249 validation. Fixed with the training median. I did not use validation rows to compute any imputation value.',
    'quote_signal, distance, equipment, dates: complete, no fix needed. No zero or negative distances.',
    'Outliers: a small tail above $12,000 (0.1%) up to $25,533. I kept them - they are long-haul Reefer loads, valid but rare. This is why RMSE (~$640) is much bigger than MAE (~$130): RMSE squares the big misses.',
])
p('Rule I followed everywhere: every cleaning number (medians, means, encodings) is computed on the training part of each fold only, then applied to validation. Nothing is ever computed on validation targets.')

h('4. EDA - key findings')
bullets([
    'Price grows almost linearly with distance. Distance is by far the strongest signal.',
    'Weight alone is nearly flat against price, but weight-per-mile and distance x weight matter (heavy + short is a hard load).',
    'Equipment matters: Reefer avg $2,553 > Flatbed $2,445 > Dry Van $2,271. Reefer (refrigerated) costs most.',
    'market_index (fuel/market level) and quote_signal (demand for that load) both push prices up, especially in combination (market x quote) and on long trips (distance x market, distance x quote).',
    'Lexington to Fort Wayne appears 32 times in training, so the December lane is not completely unseen.',
])

h('5. Feature engineering (28 features)')
p('Everything here is business logic, no target involved, so it is leakage-free:')
bullets([
    'Geography: lat_diff, lon_diff, haversine distance (R = 3958.8 miles, standard formula), road/haversine ratio and difference. The gap between road distance and straight-line distance captures detours and terrain.',
    'Time from date: day_of_week, day_of_month, month, quarter, week_of_year, is_weekend. Trees use these to learn weekday and seasonal effects - this is the only thing that varies in the December probe.',
    'Weight/distance: log_distance, log_weight (compress the long tail), weight_per_mile, distance x weight (total work). Divisions clip distance at 1 to avoid divide-by-zero.',
    'Market interactions: distance x market, market x quote, distance x quote.',
    'Categorical: route built from pickup and delivery, route_popularity (count, train-only), smoothed target encodings pickup_enc / delivery_enc / route_enc (see next section), equipment one-hot (2 columns). Raw lat/lon, city names, route strings, date, and load_id are dropped after encoding.',
])

h('6. Target encoding - leakage-safe (important)')
p('The old notebook computed city/route average rates on the full 48,000 rows before splitting. That leaks future prices into training and makes validation look better than it is (old CatBoost MAE $122.71 was optimistic). I fixed it:')
bullets([
    'Per fold: group means and counts are computed on the training part only, then mapped onto validation. Unseen cities/routes get the training global mean.',
    'Smoothing: (count x category_mean + 100 x global_mean) / (count + 100). Smoothing 100 means a route with 2 loads barely moves the global mean, while a route with 500 loads keeps its own average. This stops rare lanes from overfitting.',
    'Final model: encodings recomputed on all 48,000 rows, applied to the 12,000 validation rows the same way.',
])

h('7. Validation and split approach')
p('The real test is future data (November/December), so random splitting would cheat - it would let October predict September. I used temporal splits:')
bullets([
    'PRIMARY: train Jan 1 - Aug 31 (38,477 rows) -> validate Sep 1 - Oct 31 (9,523 rows). This is the model-selection fold.',
    'ROLL1: train Jan - Jun (28,806) -> validate July (4,912).',
    'ROLL2: train Jan - Jul (33,718) -> validate August (4,759).',
    'Every fold refits cleaning, features, and encodings on train only. Metrics are MAE (primary, in dollars - average miss), RMSE (punishes big misses), R-squared (variance explained).',
])
p('Why this way: it mimics the actual task (past predicts future), respects how market/quote/route prices drift over time, and the rolling folds prove the winner is stable, not lucky on one window.')

h2('Model comparison (primary fold)')
tbl = d.add_table(rows=1, cols=5)
tbl.style = 'Light Grid Accent 1'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, x in enumerate(['Model', 'Target', 'MAE ($)', 'RMSE ($)', 'R2']):
    hdr[i].text = x
rows = [
    ('CatBoost', 'log1p', '130.30', '639.63', '0.8243'),
    ('CatBoost', 'raw', '134.08', '638.29', '0.8251'),
    ('LightGBM', 'log1p', '136.09', '642.49', '0.8227'),
    ('LightGBM', 'raw', '148.14', '647.70', '0.8199'),
    ('XGBoost', 'log1p', '151.85', '652.01', '0.8175'),
    ('XGBoost', 'raw', '164.87', '661.44', '0.8121'),
    ('Base $/mile', '-', '256.95', '684.25', '0.7990'),
    ('Base median', '-', '1148.92', '1569.42', '-0.058'),
]
for r in rows:
    c = tbl.add_row().cells
    for i, x in enumerate(r):
        c[i].text = x
p('Rolling (raw target): ROLL1 July - CatBoost $154.25 / LightGBM $167.67 / XGBoost $240.77. ROLL2 August - CatBoost $143.53 / LightGBM $177.42 / XGBoost $256.56. CatBoost is the most stable; XGBoost overfits the early months.')
p('Target experiment: log1p(target) with expm1 back-conversion beats raw on MAE for all three models (evaluated on the original dollar scale), so the winner uses log1p. Baselines show the ML models add real value: median/mean baselines miss by about $1,150, distance-only by $257.')

h('8. Why CatBoost won')
bullets([
    'Lowest primary MAE ($130.30 with log1p) and best rolling stability.',
    'Handles the skewed target and categorical/encoded mix well without extra tuning.',
    'Hyperparameters (kept modest on purpose): iterations=1500, learning_rate=0.03, depth=7, RMSE loss, seed 42, early stopping 50. LightGBM (leaves 63, depth 7) and XGBoost (depth 7, same schedule; early_stopping_rounds in the constructor for v3.x) used comparable budgets.',
])

h('9. Feature importance (best model)')
p('Top 15 by CatBoost importance: distance (21.2), log_distance (17.2), haversine_dist (13.8), distance_x_market (8.3), route_enc (7.6), distance_x_quote (7.5), lon_diff (5.6), distance_x_weight (3.7), road_hav_diff (3.4), equipment_Reefer (2.7), weight_per_mile (1.3), quote_signal (1.1), equipment_Flatbed (1.0), route_popularity (0.7), week_of_year (0.6). Reading: price is distance first, market/demand interaction second, lane premium (route_enc) third, truck type fourth. The encodings are disclosed here - they carry city/lane price levels.')
d.add_picture('feature_importance.png', width=Inches(5.8))
d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

h('10. December fixed-lane predictions')
p('Problem: the December file has no lat/lon, market_index, or quote_signal - only city names, distance, equipment, weight, date. I did not invent flat placeholders. Method: real Lexington pickup (36.99152, -84.99876) and Fort Wayne delivery (41.31561, -85.36206) medians from training; market_index 0.9281 and Dry-Van quote_signal 2.0538 from the Sep-Oct trailing median (no December market is observed in training, so last-60-day median is the defensible forward fill); route/city encodings from the full 48k training; date features vary naturally. Range $809.61-$833.66 - gentle weekly movement with a small year-end dip, which is what the date-only variation should produce. The old notebook jump ($799 to $1,810 on Dec 26) is gone.')
d.add_picture('scorer_results/candidate_december.png', width=Inches(6.0))
d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph('Figure: scorer_results/candidate_december.png from score.py (31 days, Lexington to Fort Wayne, 360 mi, Dry Van, 32,000 lb).').italic = True

h('11. Submission files and reproducibility')
bullets([
    'validation_predictions.csv: 12,000 rows, load_id,predicted_rate, IDs match template, no duplicates/NaN, all positive. Range $163.90-$7,912.33, mean $2,336.21 (vs train mean $2,373.98 - sensible).',
    'Data/december-chart-inputs.csv: 31 rows, 7 original columns, fixed lane values untouched, predicted_rate filled.',
    'scorer_results/candidate_december.png: generated by score.py, passes validation.',
    'Run: pip install -r requirements.txt (plus lightgbm xgboost catboost scikit-learn seaborn), then python train.py, then python score.py --predictions validation_predictions.csv --december-predictions Data/december-chart-inputs.csv. Seeds fixed, relative paths, no notebook state needed. Code: train.py; notebook kept for exploration.',
])

h('12. Limitations and next steps')
bullets([
    'RMSE near $640 stays high because I kept the rare $15k+ loads. Clipping or a separate long-haul model could lower RMSE but would hide real risk.',
    'December market/quote is a proxy (Sep-Oct median). If Spotter provides December fuel/demand indices, plug them in directly.',
    'Next: light tuning of depth/learning-rate, quantile loss for the tail, and a distance-banded error check.',
])
d.add_paragraph('Bottom line: honest temporal validation, leakage fixed, CatBoost-log at $130 MAE (about 6% error), all submission files generated and scorer passes.').bold = True
d.save('Freight_Rate_Prediction_Report.docx')
print('saved docx')
